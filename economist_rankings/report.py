"""Write ranking results to a formatted Excel workbook."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TOP3_FILL = PatternFill("solid", fgColor="E2EFDA")


def _style_sheet(ws, df: pd.DataFrame, highlight_top: bool) -> None:
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"

    for idx, col in enumerate(df.columns, start=1):
        width = max(len(str(col)), *(len(str(v)) for v in df[col].head(200))) + 2
        ws.column_dimensions[get_column_letter(idx)].width = min(width, 32)

    if highlight_top and "rank" in df.columns:
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            rank_cell = row[list(df.columns).index("rank")]
            if isinstance(rank_cell.value, (int, float)) and rank_cell.value <= 3:
                for cell in row:
                    cell.fill = TOP3_FILL


def write_report(
    path: str | Path,
    overall: pd.DataFrame,
    by_indicator: pd.DataFrame,
    by_year: pd.DataFrame,
    detail: pd.DataFrame,
    notes: list[str] | None = None,
) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    sheets: list[tuple[str, pd.DataFrame, bool]] = [
        ("Overall Ranking", overall, True),
        ("By Indicator", by_indicator, False),
        ("By Year", by_year, False),
        ("Forecast Detail", detail.round(3), False),
    ]

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, df, highlight in sheets:
            df.to_excel(writer, sheet_name=name, index=False)
            _style_sheet(writer.sheets[name], df, highlight)

        if notes:
            notes_df = pd.DataFrame({"notes": notes})
            notes_df.to_excel(writer, sheet_name="Notes", index=False)
            ws = writer.sheets["Notes"]
            ws.column_dimensions["A"].width = 110
            _style_sheet(ws, notes_df, False)

    return path
